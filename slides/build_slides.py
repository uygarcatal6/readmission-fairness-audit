"""
slides/build_slides.py — the 10-minute exam talk, built from the saved results.

    python slides/build_slides.py          # writes slides/readmission_fairness_talk.pptx
    python slides/build_slides.py --pdf    # also exports the PDF through PowerPoint (Windows)

Every number on a slide and in the speaker notes is read from outputs/approved/*.json (the
same files as the paper and the notebook); nothing is typed by hand. Figures come from
readmission/figures.py, rendered at slide resolution into slides/img/. The source line at the
bottom of each result slide names the report section (outputs/reports/approved_*.md) that
prints the same numbers.

Design: 16:9, white background, one accent colour (Okabe–Ito blue, as in the figures), titles
that state the slide's message, body text >= 18 pt, speaker notes of about one minute each.
"""

from __future__ import annotations

import argparse
import math
import os
import re
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path

os.environ.setdefault("OMP_NUM_THREADS", "2")

import matplotlib  # noqa: E402

matplotlib.use("Agg")

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from readmission import figures as F  # noqa: E402

OUT_PPTX = HERE / "readmission_fairness_talk.pptx"
OUT_PDF = HERE / "readmission_fairness_talk.pdf"
IMG = HERE / "img"
FIG_DPI = 300
SLIDE_FIGURES = ("models", "fairness", "sources", "mitigation", "ordering")

# Author shown on the title slide.
AUTHOR = "Uygar Çatal"
COURSE = "Machine Learning and Data Mining · University of Bologna · A.Y. 2026/27"

ACCENT = RGBColor.from_string(F.BLUE.lstrip("#"))  # Okabe–Ito blue, as in the figures
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5F, 0x5F, 0x5F)
PALE = RGBColor(0xE8, 0xF1, 0xF8)  # light tint of the accent, for boxes
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
TITLE_PT, BODY_PT, SMALL_PT = 30, 20, 12

MINUS = "−"
DAG = "†"  # an addition beyond the approved proposal, as \add in paper/main.tex
DAG_NOTE = f"{DAG} = addition beyond the approved proposal"
NUMBER_WORDS = {2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight"}
TITLE = "Who gets flagged?"

# Figures are drawn for a 3.5-in paper column. For the slides (only here: readmission/figures.py
# and the paper figures are left unchanged) each is resized to the aspect of its box on the slide
# divided by FIG_MAG, its text and markers are enlarged, numeric axes get fewer ticks and legends
# move outside the axes, so that legends reach about 14 pt and tick labels about 16 pt on the slide.
FIG_MAG = 1.9            # slide size / figure size
FIG_TEXT_SCALE = 1.35    # 5.5-pt legends x 1.35 x 1.9 = 14 pt
FIG_MARKER_SCALE = 1.4
FIG_BOX = {              # (width, height) in inches of each figure's box on its slide
    "models": (7.6, 4.85), "fairness": (6.9, 5.3), "sources": (7.2, 5.3),
    "mitigation": (6.1, 4.85), "ordering": (7.3, 5.3)}


# ── numbers: formatted from the JSONs ─────────────────────────────────────────────────────


def n3(x: float, d: int = 3, sign: bool = False) -> str:
    """Round to d decimals with a typographic minus (and a + when sign=True)."""
    s = f"{x:+.{d}f}" if sign else f"{x:.{d}f}"
    return s.replace("-", MINUS)


def ci(c, d: int = 3) -> str:
    """A confidence interval; the no-break space keeps it on one line of a slide."""
    return f"[{n3(c[0], d)}, {n3(c[1], d)}]"


def pct(x: float, d: int = 1) -> str:
    return f"{100 * x:.{d}f} %"


def thousands(n: int) -> str:
    return f"{n:,}"


def facts(R: dict) -> dict:
    """Every number the slides print, read from the five result JSONs (with its report section)."""
    s, f, m, so, o = (R[k] for k in F.RESULT_FILES)
    v: dict = {}
    pm = f["primary_model"]                         # rf
    lead = f["leading_attribute"]                   # age_band
    key = f"{pm}__topq"
    assert lead == m["leading_attribute"] == so["leading_attribute"] == "age_band"
    v["pm"], v["lead"] = pm, lead

    # main §1–§3 (summary.json)
    v["n_enc"] = thousands(s["n_encounters"])
    v["n_pat"] = thousands(o["all_encounters"]["hba1c_requested"]["n_patients"])  # ordering §2
    n_pos = sum(g["n_pos"] for g in f["groups"]["sex"].values())                  # fairness §1
    assert sum(g["n"] for g in f["groups"]["sex"].values()) == s["n_encounters"]
    v["n_pos"] = thousands(n_pos)
    v["prev"] = n3(s["prevalence"], 3)
    v["q"] = n3(s["prevalence"], 3)
    v["q_pct"] = f"{100 * s['prevalence']:.0f}"
    v["nir"] = n3(s["nir"], 3)
    v["n_rep"] = s["n_repeats"]
    v["n_folds"] = so["n_splits"] // so["n_repeats"]
    v["n_boot"] = thousands(s["n_boot"])
    v["n_perm"] = f["n_permutations"]
    v["n_blocks"] = len(so["blocks"])
    v["n_inputs"] = sum(len(b) for b in so["blocks"].values())

    rk = s["ranking"]
    auc = {mm: rk[f"model:{mm}"]["roc_auc"] for mm in F.APPROVED}
    best = max(F.APPROVED, key=lambda mm: auc[mm]["estimate"])
    assert best == pm, "primary model = highest ROC-AUC (approved rule)"
    second = sorted(F.APPROVED, key=lambda mm: auc[mm]["estimate"])[-2]
    v["auc"] = n3(auc[pm]["estimate"])
    v["auc_ci"] = ci(auc[pm]["ci"])
    v["auc2_model"], v["auc2"] = second, n3(auc[second]["estimate"])
    pr = rk[f"model:{pm}"]["pr_auc"]["estimate"]
    v["pr_auc"] = n3(pr)
    v["pr_ratio"] = f"{pr / s['prevalence']:.1f}"
    d = s["decision"][key]
    v["tpr"] = n3(d["tpr"]["estimate"])
    v["ppv"] = n3(d["ppv"]["estimate"])
    v["accnir"] = n3(d["acc_minus_nir"]["estimate"], 4)
    v["accnir_ci"] = ci(d["acc_minus_nir"]["ci"], 4)
    v["accnir_s"] = n3(d["acc_minus_nir"]["estimate"], 3)  # 3 decimals on the slide
    v["accnir_ci_s"] = ci(d["acc_minus_nir"]["ci"], 3)
    ppv = d["ppv"]["estimate"]
    v["one_in"] = NUMBER_WORDS[round(1 / ppv)]           # "one flag in four"
    assert abs(1 / ppv - round(1 / ppv)) < 0.25, "PPV is close to one in N"
    # the gate: approved models (as in approved_main.md §3) and the course extensions
    rules = {k: x for k, x in s["decision"].items() if "usable" in x}
    appr = {k: x for k, x in rules.items() if k.split("__")[0] in F.APPROVED}
    ext = {k: x for k, x in rules.items() if k.split("__")[0] in F.EXTENSIONS}
    assert len(appr) + len(ext) == len(rules)
    v["n_rules"] = len(appr)
    v["n_usable"] = sum(bool(x["usable"]) for x in appr.values())
    v["n_models"], v["n_ext"] = NUMBER_WORDS[len(F.APPROVED)], NUMBER_WORDS[len(F.EXTENSIONS)]
    assert v["n_usable"] == 0, "slides 1, 4, 9, 10 say no model passes the gate"
    assert not any(x["usable"] for x in ext.values()), "slide 4: the extensions fail too"
    d05 = s["decision"][f"{pm}__q05"]
    v["ppv05"] = n3(d05["ppv"]["estimate"])
    v["accnir05"] = n3(d05["acc_minus_nir"]["estimate"], 4)
    assert d05["ppv"]["estimate"] < 0.5
    assert all(x["acc_minus_nir"]["ci"][1] <= d05["acc_minus_nir"]["ci"][1]
               for x in s["decision"].values() if "acc_minus_nir" in x), "q05 is the closest rule"
    qs = sorted(int(m_.group(1)) for k in s["decision"] if (m_ := re.search(r"__q(\d+)$", k)))
    v["q_min"], v["q_max"] = qs[0], qs[-1]               # capacity curve already computed

    # fairness §1–§4
    g = f["gaps"][key]
    eo = g[lead]["eo_diff"]
    v["eo"], v["eo_ci"] = n3(eo["estimate"]), ci(eo["ci"])
    v["null"] = n3(g[lead]["null"]["mean"])
    v["p"] = n3(g[lead]["null"]["p_value"])
    assert abs(g[lead]["null"]["p_value"] - 1 / (f["n_permutations"] + 1)) < 1e-9, \
        "slide 5: p is the minimum for n_perm shuffles"
    for a in ("sex", "race"):
        v[f"eo_{a}"] = n3(g[a]["eo_diff"]["estimate"])
        v[f"null_{a}"] = n3(g[a]["null"]["mean"])
        v[f"p_{a}"] = n3(g[a]["null"]["p_value"], 2)
        assert g[a]["null"]["p_value"] > 0.05, f"slide 5: {a} is within chance"
    rg = g["race"]
    assert rg["not_a_group"] == ["Unknown"]
    v["eo_race_missing"] = n3(rg["eo_diff_incl_missing"]["estimate"])
    pt = rg["pair_tpr"]
    v["pair"] = " − ".join(pt["groups"])
    v["pair_tpr"], v["pair_tpr_ci"] = n3(pt["estimate"]), ci(pt["ci"])
    ages = list(f["groups"][lead])
    young, old = ages[0], ages[-1]
    v["young"], v["old"] = young.replace("-", "–"), old
    gm = {a: f["groups"][lead][a]["metrics"][key] for a in ages}
    v["tpr_y"], v["tpr_o"] = n3(gm[young]["tpr"]["estimate"]), n3(gm[old]["tpr"]["estimate"])
    sel = {a: gm[a]["selection_rate"]["estimate"] for a in ages}
    v["sel_y"] = n3(sel[young])
    v["sel_rest"] = f"{n3(min(sel[a] for a in ages[1:]))}–{n3(max(sel[a] for a in ages[1:]))}"
    br = [f["groups"][lead][a]["base_rate"]["estimate"] for a in ages]
    v["br"] = f"{n3(min(br))}–{n3(max(br))}"
    citl = max(abs(c["citl"]["estimate"]) for c in f["calibration"][pm][lead].values())
    v["citl"] = f"observed − predicted within ±{math.ceil(citl * 1000) / 1000:.3f}"
    v["eo_logreg"] = n3(f["gaps"]["logreg__topq"][lead]["eo_diff"]["estimate"])
    v["eo_knn"] = n3(f["gaps"]["knn__topq"][lead]["eo_diff"]["estimate"])
    for mm in ("logreg", "knn"):   # "about half of it"
        assert 0.4 < f["gaps"][f"{mm}__topq"][lead]["eo_diff"]["estimate"] / eo["estimate"] < 0.6

    # sources §1–§5
    e = so["effects"]
    v["src_noprior"] = n3(e["without_prior_use@topq"][f"eo_{lead}"]["mitigated"])
    v["src_noprior_rand"] = n3(so["degradation"]["without_prior_use"][f"eo_{lead}"])
    topq_arms = {k: x[f"eo_{lead}"]["difference"] for k, x in e.items() if k.endswith("@topq")}
    assert min(topq_arms, key=topq_arms.get) == "without_prior_use@topq", "prior use moves most"
    assert (e["without_prior_use@topq"][f"eo_{lead}"]["mitigated"]
            < so["degradation"]["without_prior_use"][f"eo_{lead}_min"]), "beyond the random ref."
    v["src_noprior_dauc"] = n3(e["without_prior_use@topq"]["roc_auc"]["difference"])
    gr = so["group_rates"]["without_prior_use__topq"]
    v["src_tpr_y"], v["src_tpr_o"] = n3(gr[young]["tpr"]), n3(gr[old]["tpr"])
    v["src_noage"] = n3(e["no_age@topq"][f"eo_{lead}"]["mitigated"])
    v["src_noage_rand"] = n3(so["degradation"]["no_age"][f"eo_{lead}"])
    v["proxy"] = n3(so["proxy"][lead]["macro"]["estimate"])
    rkg = so["ranking"][lead]
    v["auc_y"], v["auc_o"] = n3(rkg[young]["roc_auc"]["estimate"]), n3(rkg[old]["roc_auc"]["estimate"])
    v["tpr_fpr_y"], v["tpr_fpr_o"] = n3(rkg[young]["tpr_at_overall_fpr"]), n3(rkg[old]["tpr_at_overall_fpr"])
    v["cut_y"] = n3(rkg[young]["tpr"] - rkg[young]["tpr_at_overall_fpr"], sign=True)
    assert rkg[young]["tpr"] > rkg[young]["tpr_at_overall_fpr"], "the shared cut-off adds to the young"
    assert (rkg[young]["tpr_at_overall_fpr"] - rkg[old]["tpr_at_overall_fpr"]
            > eo["estimate"] / 2), "slide 6: ranking quality is most of the gap"
    assert rkg[young]["roc_auc"]["estimate"] > rkg[old]["roc_auc"]["estimate"]

    # mitigation §1–§3, §5.1, §7
    me = m["effects"]
    ea = f"eo_{lead}"
    v["n_mitig"] = NUMBER_WORDS[len(F.MITIGATION_ARMS)]
    kc = me[f"rw_kc_{lead}@topq"][ea]
    v["kc"] = n3(kc["mitigated"])
    assert kc["ci"][0] > 0, "slide 7: K&C widens the gap"
    v["cell"] = n3(me[f"rw_cell_{lead}@topq"][ea]["mitigated"])
    v["cell_d"] = n3(me[f"rw_cell_{lead}@topq"][ea]["difference"], sign=True)
    v["cell_race_d"] = n3(me["rw_cell_race@topq"][ea]["difference"], sign=True)
    assert me["rw_cell_race@topq"][ea]["difference"] < me[f"rw_cell_{lead}@topq"][ea]["difference"]
    ec = me[f"eocap_{lead}"]
    v["eocap_from"], v["eocap"] = n3(ec[ea]["counterpart"]), n3(ec[ea]["mitigated"])
    v["eocap_dacc"] = n3(ec["accuracy"]["difference"], 4)
    to = me[f"to_{lead}"]
    assert to["counterpart"] == "val__youden" and ec["counterpart"].startswith("eocaptwin_")
    v["to_from"], v["to"] = n3(to[ea]["counterpart"]), n3(to[ea]["mitigated"])
    sel = m["audit"]["overall_selection"]
    v["to_flag"] = n3(sel[f"to_{lead}__eo"])
    v["to_flag_pct"] = f"{50 * (sel[f'to_{lead}__eo'] + sel['val__youden']):.0f}"  # both arms
    assert abs(sel[f"to_{lead}__eo"] - sel["val__youden"]) < 0.01, "TO and its twin flag as many"
    mg = {a: m["audit"]["groups"][lead][a]["metrics"] for a in ages}
    t = lambda a, k: mg[a][k]["tpr"]["estimate"]  # noqa: E731
    v["m_twin_y"], v["m_eo_y"] = n3(t(young, "val__topq")), n3(t(young, f"eocap_{lead}__eo"))
    v["m_twin_o"], v["m_eo_o"] = n3(t(old, "val__topq")), n3(t(old, f"eocap_{lead}__eo"))
    v["m_yd_y"], v["m_to_y"] = n3(t(young, "val__youden")), n3(t(young, f"to_{lead}__eo"))
    for twin in ("val__topq", ec["counterpart"]):   # the 80 % model at q %, and its same-number twin
        assert all(t(a, f"eocap_{lead}__eo") <= t(a, twin) for a in ages), "no TPR rises (EO cap)"
    assert all(t(a, f"to_{lead}__eo") < t(a, "val__youden") for a in ages), "every band falls (TO)"
    v["m_usable"] = sum(bool(x.get("usable")) for x in m["summary"]["decision"].values())
    assert v["m_usable"] == 0, "no mitigation passes the gate"
    v["sex_max_d"] = n3(max(abs(me[k]["eo_sex"]["difference"]) for k in me
                            if k.endswith("sex") or k.endswith("_sex@topq")))

    # ordering §2, §4
    h = o["all_encounters"]["hba1c_requested"]
    v["hba1c_rate"] = n3(h["rate"])
    v["hba1c_or_y"] = n3(h["terms"][f"age_band={young}"]["odds_ratio"], 2)
    gl, ga = o["all_encounters"]["glucose_requested"]["terms"], o["source_adjusted"]["glucose_requested"]["terms"]
    for short, lvl in (("aa", "AfricanAmerican"), ("hi", "Hispanic")):
        v[f"{short}_or"] = n3(gl[f"race={lvl}"]["odds_ratio"], 2)
        v[f"{short}_or_adj"] = n3(ga[f"race={lvl}"]["odds_ratio"], 2)
        v[f"{short}_ci_adj"] = ci(ga[f"race={lvl}"]["ci"], 2)
    hi_ci, aa_ci = ga["race=Hispanic"]["ci"], ga["race=AfricanAmerican"]["ci"]
    assert hi_ci[0] < 1 < hi_ci[1], "slide 8: the Hispanic excess vanishes"
    assert aa_ci[1] < 1, "slide 8: the African American deficit remains"
    assert gl["race=Hispanic"]["odds_ratio"] > 1 > gl["race=AfricanAmerican"]["odds_ratio"]
    rec = o["recording"]["glucose_requested"]
    v["glu_share"] = pct(rec["share_of_events_source"])
    v["glu_enc"] = pct(rec["share_of_encounters_source"])
    v["glu_rate_src"], v["glu_rate_oth"] = n3(rec["rate_source"]), n3(rec["rate_other"])
    return v


# ── figures ───────────────────────────────────────────────────────────────────────────────


def _for_slide(name: str, fig) -> None:
    """Adapt one paper figure to its slide box (see FIG_MAG); the data drawn are unchanged."""
    from matplotlib.lines import Line2D
    from matplotlib.text import Text
    from matplotlib.ticker import AutoLocator, MaxNLocator

    if name == "mitigation":
        # the slide shows only the right panel (TPR by age band per arm); the left panel's
        # Δ EO / Δ accuracy values are the bullets of slide 7
        ax_t, ax_r = fig.axes[:2]
        ax_t.remove()
        ax_r.set_subplotspec(fig.add_gridspec(1, 1)[0])
    for ax in fig.axes:
        loc = ax.xaxis.get_major_locator()   # at most 3 bins on numeric axes (fixed ticks kept)
        if isinstance(loc, AutoLocator) or (type(loc) is MaxNLocator and loc._nbins > 3):
            ax.xaxis.set_major_locator(MaxNLocator(3))
        if len(ax.get_xlabel()) > 12:               # a long label stays inside the figure
            ax.set_xlabel(ax.get_xlabel(), loc="right")
        if (leg := ax.get_legend()) is not None:   # legends go outside the axes, in the figure
            leg.remove()
            handles, labels = ax.get_legend_handles_labels()   # error-bar markers, not lines
            where, cols = ("outside right center", 1) if name == "mitigation" else                 ("outside lower center", len(labels))
            fig.legend(handles, labels, loc=where, ncols=cols, handletextpad=0.2,
                       columnspacing=0.8, labelspacing=0.3, fontsize=5.5)
    for t in fig.findobj(Text):
        t.set_fontsize(t.get_fontsize() * FIG_TEXT_SCALE)
    for ln in fig.findobj(Line2D):   # data markers and legend handles alike
        ln.set_markersize(ln.get_markersize() * FIG_MARKER_SCALE)
    w, h = FIG_BOX[name]
    fig.set_size_inches(w / FIG_MAG, h / FIG_MAG)


def render_figures(R: dict) -> dict[str, Path]:
    """The figures of readmission/figures.py, adapted to the slides, as PNGs in slides/img/."""
    IMG.mkdir(exist_ok=True)
    out = {}
    for name, fig in F.all_figures(R).items():
        if name not in SLIDE_FIGURES:
            continue
        _for_slide(name, fig)
        path = IMG / f"{name}.png"
        fig.savefig(path, dpi=FIG_DPI)
        out[name] = path
    return out


# ── slide helpers ─────────────────────────────────────────────────────────────────────────


def _font(run, size, *, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _runs(p, text: str, size: int, color=INK, bold=False):
    """Add text to paragraph p; **x** is printed bold in the accent colour."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        emph = i % 2 == 1
        _font(r, size, bold=bold or emph, color=ACCENT if emph else color)


def _bullet(p, indent_in: float = 0.32):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent_in))))
    pPr.set("indent", str(-Emu(Inches(indent_in))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", str(ACCENT))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def textbox(slide, left, top, width, height, paras, *, size=BODY_PT, bullets=False,
            color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=10):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = 1.0
        _runs(p, text, size, color=color, bold=bold)
        if bullets:
            _bullet(p)
    return tb


def box(slide, left, top, width, height, *, fill=PALE, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(2)
    return s


def boxed_text(slide, left, top, width, height, paras, *, fill=PALE, line=None, size=BODY_PT,
               color=INK, bold=False, align=PP_ALIGN.LEFT, shape=MSO_SHAPE.RECTANGLE,
               anchor=MSO_ANCHOR.MIDDLE, pad=0.15):
    box(slide, left, top, width, height, fill=fill, line=line, shape=shape)
    p = Inches(pad)
    return textbox(slide, left + p, top + p, width - 2 * p, height - 2 * p, paras, size=size,
                   color=color, bold=bold, align=align, anchor=anchor, space=4)


def picture(slide, path: Path, left, top, max_w, max_h, *, align="center"):
    """Add a PNG scaled to fit (max_w, max_h), keeping its aspect ratio."""
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    x = left + (max_w - pw) // 2 if align == "center" else left
    y = top + (max_h - ph) // 2
    return slide.shapes.add_picture(str(path), x, y, pw, ph)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.notes: list[str] = []
        # replace the python-pptx template's metadata (its author, dates and description)
        cp = self.prs.core_properties
        cp.author = cp.last_modified_by = AUTHOR
        cp.title = TITLE
        cp.subject = "A CRISP-DM fairness audit of 30-day readmission prediction"
        cp.comments = "Exam talk, built by slides/build_slides.py from outputs/approved/*.json"  # dc:description
        cp.keywords = cp.category = ""
        cp.created = cp.modified = datetime.now(UTC).replace(microsecond=0, tzinfo=None)
        cp.revision = 1
        for part in self.prs.part.package.iter_parts():
            if str(part.partname) == "/docProps/app.xml":
                part._blob = part.blob.replace(b"On-screen Show (4:3)", b"Widescreen")

    def slide(self, title: str | None, notes: str, source: str | None = None):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = WHITE
        if title:
            textbox(s, MARGIN, Inches(0.35), W - 2 * MARGIN, Inches(0.95), [title],
                    size=TITLE_PT, bold=True, anchor=MSO_ANCHOR.BOTTOM, space=0)
            box(s, MARGIN, Inches(1.36), Inches(1.3), Inches(0.07), fill=ACCENT)
        if source:
            textbox(s, MARGIN, Inches(7.0), Inches(10.5), Inches(0.35), ["Source: " + source],
                    size=SMALL_PT, color=MUTED, space=0)
        s.notes_slide.notes_text_frame.text = " ".join(notes.split())
        self.notes.append(" ".join(notes.split()))
        return s

    def number_slides(self):
        total = len(self.prs.slides)
        for i, s in enumerate(self.prs.slides, start=1):
            if i == 1:
                continue
            textbox(s, W - MARGIN - Inches(1.2), Inches(7.0), Inches(1.2), Inches(0.35),
                    [f"{i} / {total}"], size=SMALL_PT, color=MUTED, align=PP_ALIGN.RIGHT, space=0)


def words(text: str) -> int:
    return len(re.findall(r"\S+", text))


# ── the talk ──────────────────────────────────────────────────────────────────────────────

FIG_TOP = Inches(1.6)
FIG_H = Inches(5.3)


def figure_slide(deck, title, fig: Path, bullets, notes, source, *, fig_w=7.5, callout=None,
                 caption=None):
    s = deck.slide(title, notes, source + " · " + DAG_NOTE)
    fw = Inches(fig_w)
    fh = FIG_H - (Inches(0.45) if caption else 0)
    picture(s, fig, MARGIN - Inches(0.1), FIG_TOP, fw, fh)
    if caption:
        textbox(s, MARGIN, FIG_TOP + fh + Inches(0.02), fw, Inches(0.4), [caption],
                size=16, color=MUTED, align=PP_ALIGN.CENTER, space=0)
    tx = MARGIN + fw + Inches(0.2)
    tw = W - MARGIN - tx
    th = FIG_H - (Inches(1.25) if callout else 0)
    textbox(s, tx, FIG_TOP + Inches(0.05), tw, th, bullets, bullets=True, space=12)
    if callout:
        boxed_text(s, tx, FIG_TOP + FIG_H - Inches(1.1), tw, Inches(1.05), [callout],
                   fill=ACCENT, color=WHITE, bold=True, size=BODY_PT, align=PP_ALIGN.CENTER)
    return s


def count_tests() -> str:
    """Test cases pytest collects for the approved protocol (parametrised cases count once each)."""
    files = sorted(str(p.relative_to(ROOT)) for p in (ROOT / "tests").glob("test_readmission_*.py"))
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    r = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q", "-p",
                        "no:cacheprovider", *files], cwd=ROOT, env=env, capture_output=True,
                       text=True, check=False)
    m = re.search(r"(\d+) tests? collected", r.stdout)
    assert m, f"pytest collection failed:\n{r.stdout[-2000:]}\n{r.stderr[-2000:]}"
    return m.group(1)


def build(R: dict, img: dict[str, Path]) -> Deck:
    v = facts(R)
    d = Deck()

    # 1 — title
    s = d.slide(None, """
        Good morning. A hospital can follow up only a few of its discharged diabetic patients
        and wants a model to choose who. I asked two questions: is any model usable at that
        fixed capacity, and are its errors shared equally across sex, age band and race? The
        answers up front: no model passes the usability gate. The clearest disparity is by age.
        The two post-processors that close it do so by lowering true-positive rates: levelling
        down, not a fix. And the race differences in recorded glucose results are largely, not
        wholly, a recording pattern of the data source.""")
    box(s, 0, 0, Inches(0.35), H, fill=ACCENT)
    textbox(s, Inches(1.1), Inches(1.7), Inches(11.2), Inches(1.2), [TITLE],
            size=48, bold=True, anchor=MSO_ANCHOR.BOTTOM, space=0)
    textbox(s, Inches(1.1), Inches(2.95), Inches(11.2), Inches(1.2),
            ["A CRISP-DM fairness audit of 30-day readmission prediction",
             "UCI Diabetes 130-US Hospitals"], size=26, color=ACCENT, space=6)
    box(s, Inches(1.1), Inches(4.5), Inches(1.3), Inches(0.07), fill=ACCENT)
    textbox(s, Inches(1.1), Inches(4.8), Inches(11.2), Inches(1.4),
            [AUTHOR, COURSE, "Individual assignment · Option 2 (CRISP-DM)"], size=20,
            color=MUTED, space=6)

    # 2 — problem and capacity rule
    s = d.slide("The question: who gets the few follow-up slots?", f"""
        The hospital ranks every discharged patient by predicted risk and flags the top q
        percent, where q is the readmission rate of the training data, about {v['q']}. That fixes
        the number of follow-up slots without reading any test label. The fairness goal is
        equalized odds: every group should have the same true-positive rate among patients who
        really come back, and the same false-positive rate among those who do not. The approved proposal fixed a usability gate before the protocol
        was run: the model counts as usable only if the lower confidence bound of accuracy
        minus the no-information rate, {v['nir']}, is above zero. With about {v['q_pct']}
        percent readmitted, that means a PPV above one half: most flags must be right.""",
               "docs/APPROVED_PROPOSAL.md §2; approved_main.md header; approved_mitigation.md §3 note")
    steps = [("1  Rank", "every discharged diabetic inpatient by predicted 30-day readmission risk"),
             ("2  Flag the top q %", f"q = training-fold readmission rate (overall **{v['q']}**): "
                                     "a fixed follow-up capacity"),
             ("3  Audit the errors", "equal TPR and FPR across sex, age band and race "
                                     "(equalized odds)")]
    bw, gap, top = Inches(3.75), Inches(0.425), Inches(1.75)
    for i, (head, body) in enumerate(steps):
        x = MARGIN + i * (bw + gap)
        box(s, x, top, bw, Inches(2.45), fill=PALE)
        box(s, x, top, bw, Inches(0.62), fill=ACCENT)
        textbox(s, x + Inches(0.15), top, bw - Inches(0.3), Inches(0.62), [head], size=22,
                bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, space=0)
        textbox(s, x + Inches(0.15), top + Inches(0.75), bw - Inches(0.3), Inches(1.6), [body],
                size=BODY_PT, space=0)
        if i < 2:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.07), top + Inches(1.0),
                                   gap - Inches(0.14), Inches(0.45))
            a.fill.solid()
            a.fill.fore_color.rgb = ACCENT
            a.line.fill.background()
    boxed_text(s, MARGIN, Inches(4.55), W - 2 * MARGIN, Inches(2.1), [
        "**Usability gate** (fixed in the approved proposal, before the approved protocol was "
        f"run): usable only if the lower 95 % bound of accuracy − NIR ({v['nir']}) is > 0.",
        "With m flagged, accuracy − NIR = (2·TP − m) / n, so the gate needs **PPV > 0.5**: "
        "most flags must be right."], fill=WHITE, line=ACCENT, pad=0.25)

    # 3 — data and pipeline
    s = d.slide("Data and pipeline: CRISP-DM, validated patient by patient", f"""
        The data are the public Diabetes 130-US Hospitals file from the UCI repository. After
        removing deaths and hospice discharges, unknown sex, missing primary diagnosis and
        patients under 20, we keep {v['n_enc']} encounters from {v['n_pat']} patients;
        {v['n_pos']} of them were readmitted within 30 days. Many patients come back several times, so the folds are grouped by patient: the same
        person is never in training and test at once. We use {v['n_folds']} folds, repeated
        {v['n_rep']} times. Encoding and scaling are fitted on training rows only, tree depth
        and k are chosen by the one-standard-error rule on inner patient folds, and every
        interval comes from {v['n_boot']} bootstrap draws over patients.""",
               "approved_main.md header; approved_fairness.md §1; approved_sources.md block table")
    rows = [("Business understanding", "flag the top q % at a fixed capacity; equal errors by sex, age band, race"),
            ("Data understanding", f"**{v['n_enc']}** encounters · **{v['n_pat']}** patients · "
                                   f"**{v['n_pos']}** readmitted in 30 days ({v['prev']})"),
            ("Data preparation", f"{v['n_inputs']} inputs in {v['n_blocks']} blocks; encoding and "
                                 "scaling fitted on training rows only"),
            ("Modelling", "tree, k-NN, logistic regression, random forest, gradient boosting; "
                          "depth and k by the one-SE rule"),
            ("Evaluation", f"**patient-level** stratified {v['n_folds']}-fold CV × {v['n_rep']} "
                           f"repeats; {v['n_boot']} bootstrap draws"),
            ("Deployment", "a recommendation for the hospital, with its limits (slide 9)")]
    rh, top = Inches(0.8), Inches(1.72)
    lw = Inches(3.3)
    for i, (phase, what) in enumerate(rows):
        y = top + i * (rh + Inches(0.07))
        boxed_text(s, MARGIN, y, lw, rh, [phase], fill=ACCENT, color=WHITE, bold=True,
                   shape=MSO_SHAPE.PENTAGON, size=BODY_PT, pad=0.08)
        boxed_text(s, MARGIN + lw + Inches(0.15), y, W - 2 * MARGIN - lw - Inches(0.15), rh,
                   [what], fill=PALE, size=BODY_PT, pad=0.08)

    # 4 — models and the gate
    figure_slide(d, "Models rank better than chance, but none passes the usability gate",
                 img["models"], [
                     f"Random forest: ROC-AUC **{v['auc']}** {v['auc_ci']}",
                     f"PR-AUC **{v['pr_auc']}**, {v['pr_ratio']}× the prevalence ({v['prev']})",
                     f"At capacity: TPR = PPV = **{v['ppv']}**",
                     f"Accuracy − NIR **{v['accnir_s']}** {v['accnir_ci_s']}",
                     f"Closest rule (top 5 %): PPV {v['ppv05']}, still below 0.5"],
                 f"""
        {v['n_models'].capitalize()} approved models in black, {v['n_ext']} course extensions
        in grey. The random forest has the highest ROC-AUC, {v['auc']}, with gradient boosting
        at {v['auc2']}, inside its interval. That is clearly better than chance, and its PR-AUC
        is about {v['pr_ratio']} times the prevalence. But at the capacity rule the forest's
        true-positive rate and PPV are both {v['ppv']}: only about one flag in {v['one_in']} is
        a real readmission. Accuracy minus the no-information rate is {v['accnir']}, and the
        whole interval is below zero. So the gate fails for all {v['n_rules']} approved model
        and rule combinations, and the extensions fail too; the closest is the top five percent,
        with a PPV of {v['ppv05']}. I still audit fairness, because a hospital that ranks
        patients still distributes errors.""",
                 "approved_main.md §1–§3 · outputs/approved/summary.json",
                 callout=f"Gate fails: {v['n_usable']} of {v['n_rules']} approved model × rule "
                         "pairs pass (extensions fail too)",
                 caption=f"black = approved models, grey = course extensions{DAG}",
                 fig_w=FIG_BOX["models"][0])

    # 5 — fairness audit
    figure_slide(d, "The errors differ by age, not clearly by sex or race", img["fairness"], [
        f"Age band: EO difference **{v['eo']}** {v['eo_ci']}; permutation null{DAG} {v['null']} "
        f"(p = {v['p']}, the minimum for {v['n_perm']} shuffles)",
        f"TPR **{v['tpr_y']}** at {v['young']} vs **{v['tpr_o']}** at {v['old']}; the young are "
        f"flagged more ({v['sel_y']} vs {v['sel_rest']})",
        f"Sex {v['eo_sex']} (null {v['null_sex']}); race {v['eo_race']} (null {v['null_race']}; "
        f"“Unknown” = missing code, left out; {v['eo_race_missing']} with it): within chance",
        f"Not base rates ({v['br']}), not calibration ({v['citl']})"],
                 f"""
        This is the audit of the random forest at the capacity rule: each dot is a group's
        true-positive rate, left, and false-positive rate, right. Age band stands out. Its
        equalized-odds difference is {v['eo']}, against {v['null']} expected by chance from a
        permutation null, and no shuffle reached it. Readmitted patients aged 20 to 39 are
        flagged {v['tpr_y']} of the time, those over 80 only {v['tpr_o']}. Sex and race stay
        within their nulls; race Unknown is a missing code, left out. The
        {v['pair'].replace(' − ', ' versus ')} difference alone has an interval above zero, but
        that extreme pair is chosen after seeing the data; the permutation null allows for that
        choice. The age gap is not base rates or calibration, and logistic regression and k-NN
        show about half of it.""",
                 "approved_fairness.md §2, §3.2, §4 · outputs/approved/fairness.json",
                 fig_w=FIG_BOX["fairness"][0])

    # 6 — sources
    figure_slide(d, "The age gap comes mainly from ranking young patients better",
                 img["sources"], [
                     f"Ranked within their own group{DAG}, at the overall FPR: TPR "
                     f"**{v['tpr_fpr_y']}** ({v['young']}) vs **{v['tpr_fpr_o']}** ({v['old']}), "
                     f"most of the {v['eo']} gap; the shared cut-off adds {v['cut_y'].lstrip('+')} to the "
                     f"young (ROC-AUC {v['auc_y']} vs {v['auc_o']})",
                     f"Without prior-use counts: {v['eo']} → **{v['src_noprior']}** (same AUC "
                     f"loss at random{DAG}: {v['src_noprior_rand']}); it works by lowering "
                     f"{v['young']} TPR ({v['tpr_y']} → {v['src_tpr_y']}): no remedy",
                     f"Without age: {v['src_noage']}; other inputs predict age band "
                     f"(ROC-AUC {v['proxy']})"],
                 f"""
        Where does the age gap come from? Mainly from ranking quality. Within their own group,
        the model separates readmitted young patients with ROC-AUC {v['auc_y']}, but patients
        over 80 with only {v['auc_o']}. At the overall false-positive rate their true-positive
        rates would be {v['tpr_fpr_y']} and {v['tpr_fpr_o']}: most of the gap. The shared
        cut-off adds {v['cut_y'].lstrip('+')} to the young. I also refitted the forest without
        one block of inputs at a time, against the same ROC-AUC loss at random. Removing
        the prior-use counts moves the gap most, to {v['src_noprior']}, well beyond the random
        reference, but by lowering young patients' true-positive rate: no remedy. Removing age
        narrows it less, to {v['src_noage']}, and the other inputs still predict age band with
        ROC-AUC {v['proxy']}, so unawareness does not remove the information.""",
                 "approved_sources.md §1–§5 · outputs/approved/sources.json",
                 fig_w=FIG_BOX["sources"][0])

    # 7 — mitigations
    figure_slide(d, "The mitigations that close the gap do so by levelling down",
                 img["mitigation"], [
        f"K&C reweighing widens it: {v['eo']} → {v['kc']}",
        f"Cell balancing: {v['eo']} → {v['cell']} ({v['cell_d']}); balancing race cells cuts "
        f"it more ({v['cell_race_d']}): mostly class balancing",
        f"EO at capacity{DAG} (vs same 80 % model, same number flagged): {v['eocap_from']} → "
        f"**{v['eocap']}**; TPR {v['young']} {v['m_twin_y']} → **{v['m_eo_y']}**, {v['old']} "
        f"{v['m_twin_o']} → {v['m_eo_o']}",
        f"ThresholdOptimizer (vs one Youden threshold; both flag ≈ {v['to_flag_pct']} %): "
        f"{v['to_from']} → {v['to']}; every age band falls ({v['young']} {v['m_yd_y']} → "
        f"{v['m_to_y']})"],
                 f"""
        I tried {v['n_mitig']} mitigations, each against its unconstrained counterpart under the
        same rule. Kamiran-Calders reweighing widens the gap, to {v['kc']}. Cell balancing
        narrows it, but balancing race cells narrows the age gap even more, so it works mostly
        by balancing the classes. The two post-processors do close it. Equalized-odds
        post-processing at capacity, an addition, is compared with the same model flagging as
        many: the gap falls from {v['eocap_from']} to {v['eocap']}, but in the right panel young
        patients' true-positive rate falls from {v['m_twin_y']} to {v['m_eo_y']}, while the
        over-80s stay at about {v['m_eo_o']}. ThresholdOptimizer is compared with one Youden
        threshold; both flag about {v['to_flag_pct']} percent, the hollow markers, and every age
        band falls. Neither post-processor raises any age band's TPR against its twin: levelling
        down, not a fix.""",
                 "approved_mitigation.md §1–§3, §5.1, §7 · outputs/approved/mitigation.json",
                 callout="Neither post-processor raises any age band's TPR: levelling down, "
                         "not a fix",
                 caption=f"Filled markers flag ≈ q %, hollow ≈ {v['to_flag_pct']} %: compare "
                         "within each set", fig_w=FIG_BOX["mitigation"][0])

    # 8 — test recording
    figure_slide(d, "Glucose race gaps are largely, not wholly, a recording pattern",
                 img["ordering"], [
                     f"**{v['glu_share']}** of glucose results come from the **{v['glu_enc']}** "
                     "of encounters with unknown admission type",
                     f"With that indicator (post hoc sensitivity fit{DAG}, not tested):",
                     f"Hispanic excess vanishes: {v['hi_or']} → **{v['hi_or_adj']}** "
                     f"{v['hi_ci_adj']}",
                     f"African American deficit shrinks but remains: {v['aa_or']} → "
                     f"**{v['aa_or_adj']}** {v['aa_ci_adj']}",
                     "Odds ratios describe who has a recorded result; on their own they are not "
                     "evidence of unequal care"],
                 f"""
        Finally, measurement. The data record whether an HbA1c or glucose result exists, a
        proxy for an ordered test. A logistic model relates that to sex, age band and race,
        adjusting for the primary diagnosis. For glucose, African American patients
        have lower odds and Hispanic patients higher odds than Caucasian patients. But
        {v['glu_share'].replace(' %', ' percent')} of all glucose results come from the
        {v['glu_enc'].replace(' %', ' percent')} of encounters whose admission type is unknown,
        a recording pattern of the source. With that indicator, a sensitivity fit added after
        seeing the results, the Hispanic excess falls to {v['hi_or_adj']},
        indistinguishable from one, while the African American deficit shrinks to
        {v['aa_or_adj']} but remains. So the gap is largely, not wholly, a recording pattern,
        and on their own the odds ratios are not evidence of unequal care.""",
                 "approved_ordering.md §2, §4 · outputs/approved/ordering.json",
                 fig_w=FIG_BOX["ordering"][0])

    # 9 — what a hospital should do
    s = d.slide("What a hospital should do with this model", f"""
        First, do not use this model as a stand-alone classifier: only about
        one flag in {v['one_in']} is right, and the agreed gate needs more than one in two.
        Second, if it still allocates follow-ups by the ranking, it should report the true-positive rate by age band next to
        overall performance, because older patients are the ones the model misses. Third,
        imposing equal true-positive rates is a policy choice with a cost: here it means catching
        fewer young patients without catching more old ones. Finally, better information on
        older patients' prior care, and complete recording of admission type and lab results,
        may do more for equal errors than post-processing, although I did not test that.""",
               "paper §Conclusions (deployment); approved_main.md §2; approved_mitigation.md §5.1, §7")
    cards = [("Do not use it alone",
              f"PPV at capacity **{v['ppv']}**; the gate needs more than 0.5."),
             ("Report TPR by age band",
              f"beside overall performance: **{v['tpr_y']}** at {v['young']} vs "
              f"**{v['tpr_o']}** at {v['old']}."),
             ("Equal TPR is a policy choice",
              "at capacity it lowers young patients' TPR and raises no one's: a cost, not a fix."),
             ("Improve the data",
              "older patients' prior care; complete admission type and lab recording "
              "(not tested here).")]
    cw, ch = Inches(5.9), Inches(2.4)
    for i, (head, body) in enumerate(cards):
        x = MARGIN + (i % 2) * (cw + Inches(0.3))
        y = Inches(1.75) + (i // 2) * (ch + Inches(0.25))
        box(s, x, y, cw, ch, fill=PALE)
        box(s, x, y, Inches(0.12), ch, fill=ACCENT)
        textbox(s, x + Inches(0.35), y + Inches(0.2), cw - Inches(0.55), Inches(0.6),
                [f"{i + 1}. {head}"], size=24, bold=True, color=ACCENT, space=0)
        textbox(s, x + Inches(0.35), y + Inches(0.9), cw - Inches(0.55), ch - Inches(1.05),
                [body], size=BODY_PT, space=0)

    # 10 — limitations and next steps
    s = d.slide("Limitations, and what I would do next", f"""
        The public file does not identify the hospital or the period, so
        standard errors and the bootstrap are clustered by patient only, although the recording
        analysis shows that the source matters. The intervals are conditional on the fitted
        models. A recorded result is only a proxy for an ordered test. And the method changed:
        a new proposal was approved during the project. Open choices were fixed in writing
        before the results they affect, with one exception, the admission-type sensitivity fit.
        Next, I would validate on data that identify hospital and period, try to improve the
        ranking for older patients instead of post-processing, and agree the capacity q with
        the hospital, reporting the whole capacity curve, already computed from {v['q_min']} to
        {v['q_max']} percent. Thank you.""",
               "paper §Conclusions (limitations); docs/APPROVED_PROPOSAL.md §5")
    cols = [("Limitations", [
        "Hospital and period are not in the public file: standard errors and bootstrap "
        "clustered by patient only",
        "Intervals are conditional on the fitted models",
        "“Result recorded” is a proxy for “test ordered”",
        "Method changed mid-project (new proposal approved); admission-type fit post hoc"]),
            ("Next steps", [
        "Validate on data with hospital and period identifiers",
        "Improve the ranking for older patients (e.g. richer prior-care data) instead of "
        "post-processing",
        f"Agree q with the hospital and report the whole capacity curve "
        f"(q = {v['q_min']}–{v['q_max']} % already computed)"])]
    cw = Inches(5.9)
    for i, (head, items) in enumerate(cols):
        x = MARGIN + i * (cw + Inches(0.3))
        box(s, x, Inches(1.75), cw, Inches(0.6), fill=ACCENT)
        textbox(s, x + Inches(0.2), Inches(1.75), cw - Inches(0.4), Inches(0.6), [head], size=24,
                bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, space=0)
        textbox(s, x + Inches(0.1), Inches(2.55), cw - Inches(0.2), Inches(3.1), items,
                bullets=True, space=12)
    boxed_text(s, MARGIN, Inches(5.9), W - 2 * MARGIN, Inches(0.95), [
        "**Take-home:** the model ranks better than chance but is not usable alone; its errors "
        "differ by age; equalising them by post-processing is levelling down."],
        fill=WHITE, line=ACCENT, pad=0.2)

    # 11 — backup: reproducibility
    n_tests = count_tests()
    s = d.slide("Backup: how to reproduce every number", f"""
        This is a backup slide, in case you ask about reproducibility. The pipeline is five
        scripts run in order: the cohort and the models, then the fairness audit, the
        mitigations, the source analysis and the recording model. Each writes a report and a
        JSON summary. Every outer split is saved as a checkpoint together with a fingerprint of
        the settings, the prediction code and the cohort; a checkpoint that does not match is
        refused, never mixed with new results. A Colab notebook fits one outer split live and
        redraws every table and figure from the saved JSON files. For the approved protocol,
        pytest collects {n_tests} tests. The paper's figures and these slides are built by
        scripts from the same JSON files, so no number on a slide was typed by hand.""",
               "README.md §Run; readmission/evaluate.py (fingerprint); slides/README.md")
    textbox(s, MARGIN, Inches(1.75), Inches(12.1), Inches(5.0), [
        "Five scripts in order: **run_approved** → **run_fairness** → **run_mitigation** → "
        "**run_sources** → **run_ordering**; each writes outputs/reports/approved_*.md and "
        "outputs/approved/*.json",
        "Checkpoints per outer split carry a fingerprint of settings, prediction code and "
        "cohort; a mismatch is refused, never mixed",
        "Colab notebook: fits one outer split live and redraws every table and figure from "
        "the saved JSONs",
        f"pytest collects {n_tests} tests for the approved protocol (tests/test_readmission_*.py)",
        "Paper figures (scripts/make_figures.py) and these slides (slides/build_slides.py) "
        "read the same JSONs: no number typed by hand"], bullets=True, space=16)

    d.number_slides()
    return d


# ── export ────────────────────────────────────────────────────────────────────────────────


def export_pdf(pptx: Path, pdf: Path) -> None:
    """PowerPoint (COM) → PDF. Windows with PowerPoint only."""
    import win32com.client  # pywin32

    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(pptx), True, False, False)  # ReadOnly, Untitled, WithWindow
        try:
            pres.SaveAs(str(pdf), 32)  # ppSaveAsPDF
        finally:
            pres.Close()
    finally:
        app.Quit()


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--results", default=str(ROOT), help="repository root or results folder")
    ap.add_argument("--pdf", action="store_true", help="also export the PDF via PowerPoint")
    args = ap.parse_args(argv)
    R = F.load_results(args.results)
    img = render_figures(R)
    deck = build(R, img)
    deck.prs.save(OUT_PPTX)
    counts = [words(n) for n in deck.notes]
    for i, c in enumerate(counts, start=1):
        print(f"slide {i:2d}: {c} note words")
    main_talk = sum(counts[:10])
    print(f"notes: {main_talk} words on slides 1-10 (+{sum(counts[10:])} backup); wrote {OUT_PPTX}")
    if args.pdf:
        export_pdf(OUT_PPTX, OUT_PDF)
        print(f"wrote {OUT_PDF}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
